"""Tables for serial printing: write a ready-to-fill Excel/CSV file for a label and read filled tables.

The Excel file ptdesk writes has the label's columns in row 1 (with hints as cell comments), an
example row and the label template itself in a hidden sheet, so that opening the filled table
restores the matching label.
"""

from __future__ import annotations

import csv
import io
import json
import warnings
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from ptdesk.render import merge
from ptdesk.render.label import Label

TABLE_SUFFIXES = (".xlsx", ".xlsm", ".csv", ".txt")
DATA_SHEET = "Etiketten"
HELP_SHEET = "Hinweise"
TEMPLATE_SHEET = "ptdesk"
TEMPLATE_MARKER = "ptdesk-vorlage"
_CHUNK = 30_000  # an Excel cell holds at most 32767 characters
_VALIDATED_ROWS = 5000


class TableError(Exception):
    pass


@dataclass
class Table:
    columns: list[str]
    rows: list[dict[str, str]]
    lines: list[int]  # spreadsheet row number of every row, for messages
    path: Path | None = None
    template: dict[str, Any] | None = None  # label embedded by ptdesk (Label.to_dict format)
    template_name: str = ""
    warnings: list[str] = field(default_factory=list)


# --- writing -------------------------------------------------------------------

def template_columns(label: Label) -> tuple[list[merge.Column], bool]:
    """Columns of the label and whether the extra "Anzahl" column is added."""
    cols = merge.columns(label)
    with_copies = all(merge.key(c.name) != merge.key(merge.COPIES_COLUMN) for c in cols)
    return cols, with_copies


def write_template(path: Path, label: Label, name: str = "") -> None:
    """Write an empty table (header + example row) for the label's columns."""
    cols, _ = template_columns(label)
    if not cols:
        raise TableError("Das Etikett enthält noch keine Tabellenfelder.")
    if path.suffix.lower() == ".csv":
        _write_csv(path, label)
    else:
        _write_xlsx(path, label, name)


def _header_and_example(label: Label) -> tuple[list[str], list[str]]:
    cols, with_copies = template_columns(label)
    example = merge.example_values(label)
    headers = [c.name for c in cols]
    values = [example.get(merge.key(c.name), "") for c in cols]
    if with_copies:
        headers.append(merge.COPIES_COLUMN)
        values.append("1")
    return headers, values


def _write_csv(path: Path, label: Label) -> None:
    headers, values = _header_and_example(label)
    with path.open("w", encoding="utf-8-sig", newline="") as f:  # BOM: Excel then detects UTF-8
        writer = csv.writer(f, delimiter=";")
        writer.writerow(headers)
        writer.writerow(values)


def _write_xlsx(path: Path, label: Label, name: str) -> None:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    cols, with_copies = template_columns(label)
    example = merge.example_values(label)

    wb = Workbook()
    wb.properties.creator = "ptdesk"
    ws = wb.active
    ws.title = DATA_SHEET
    help_ws = wb.create_sheet(HELP_SHEET)

    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDEBF7")

    def header(column: int, text: str, hint: str) -> str:
        letter = get_column_letter(column)
        cell = ws.cell(row=1, column=column, value=text)
        cell.font, cell.fill = bold, header_fill
        cell.comment = Comment(hint, "ptdesk", width=300, height=120)
        return letter

    def text_cell(row: int, column: int, value: str) -> None:
        cell = ws.cell(row=row, column=column, value=value or None)
        cell.number_format = "@"
        if value.startswith("="):
            cell.data_type = "s"  # a text, not a formula

    for i, col in enumerate(cols, start=1):
        letter = header(i, col.name, col.hint)
        value = example.get(merge.key(col.name), "")
        # text format keeps "+49 …", leading zeros and dates exactly as typed
        ws.column_dimensions[letter].number_format = "@"
        ws.column_dimensions[letter].width = min(max(len(col.name), len(value), 8) + 3, 60)
        text_cell(2, i, value)

    if with_copies:
        column = len(cols) + 1
        letter = header(
            column, merge.COPIES_COLUMN,
            "Optional: wie oft das Etikett dieser Zeile gedruckt wird.\n"
            "Leer = Anzahl aus ptdesk, 0 = Zeile überspringen.",
        )
        ws.column_dimensions[letter].width = 10
        ws.cell(row=2, column=column, value=1)
        dv = DataValidation(type="whole", operator="between", formula1="0", formula2=str(merge.MAX_COPIES))
        dv.errorStyle = "warning"
        dv.error = f"Ganze Zahl von 0 bis {merge.MAX_COPIES}"
        dv.add(f"{letter}2:{letter}{_VALIDATED_ROWS}")
        ws.add_data_validation(dv)

    ws.freeze_panes = "A2"

    # help sheet
    lines = [
        "So füllst du die Tabelle aus",
        "",
        f"• Im Blatt „{DATA_SHEET}“ steht jede Zeile für ein Etikett.",
        "• Die Überschriften in Zeile 1 nicht ändern, ptdesk ordnet die Werte über sie zu.",
        "• Zeile 2 ist ein Beispiel. Überschreiben oder löschen; unverändert wird sie nicht gedruckt.",
        "• Zeilenumbrüche in einer Zelle (Alt+Enter) werden auf dem Etikett zu neuen Zeilen.",
        f"• „{merge.COPIES_COLUMN}“ (falls vorhanden): wie oft die Zeile gedruckt wird, 0 = überspringen.",
        "• Danach in ptdesk: Seriendruck → Tabelle öffnen …",
        "",
    ]
    for row, text in enumerate(lines, start=1):
        help_ws.cell(row=row, column=1, value=text)
    help_ws["A1"].font = Font(bold=True, size=13)
    start = len(lines) + 1
    for column, text in enumerate(("Spalte", "Verwendet in"), start=1):
        help_ws.cell(row=start, column=column, value=text).font = bold
    for row, col in enumerate(cols, start=start + 1):
        help_ws.cell(row=row, column=1, value=col.name)
        help_ws.cell(row=row, column=2, value=", ".join(col.usages))
    for letter, width in (("A", 28), ("B", 60)):
        help_ws.column_dimensions[letter].width = width
    for row in help_ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    # the label itself, so that opening the filled table restores it
    meta = wb.create_sheet(TEMPLATE_SHEET)
    meta.sheet_state = "veryHidden"
    meta["A1"] = TEMPLATE_MARKER
    meta["B1"] = name
    payload = json.dumps(label.to_dict(), ensure_ascii=False, separators=(",", ":"))
    for row, offset in enumerate(range(0, len(payload), _CHUNK), start=2):
        meta.cell(row=row, column=1, value="|" + payload[offset : offset + _CHUNK])  # "|": never a formula

    wb.active = 0
    wb.save(path)


# --- reading -------------------------------------------------------------------

def read_table(path: Path) -> Table:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        table = _read_xlsx(path)
    elif suffix in (".csv", ".txt", ".tsv"):
        table = _read_csv(path)
    elif suffix in (".xls", ".ods", ".numbers"):
        raise TableError("Dieses Format kann ptdesk nicht lesen. Bitte als .xlsx oder .csv speichern.")
    else:
        raise TableError("Unbekanntes Dateiformat. Unterstützt: Excel (.xlsx) und CSV.")
    table.path = path
    return table


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "ja" if value else "nein"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value).replace(".", ",")
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y" if value.time() == time() else "%d.%m.%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    return str(value)


def _read_xlsx(path: Path) -> Table:
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")  # e.g. "Data Validation extension is not supported"
            wb = load_workbook(path, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError) as e:
        raise TableError(f"Die Excel-Datei kann nicht gelesen werden: {e}") from e

    template, name = _embedded_template(wb)
    ignore = {HELP_SHEET, TEMPLATE_SHEET}
    if DATA_SHEET in wb.sheetnames:
        ws = wb[DATA_SHEET]
    else:
        ws = next((s for s in wb.worksheets if s.sheet_state == "visible" and s.title not in ignore), None)
    if ws is None:
        raise TableError("Die Datei enthält kein Tabellenblatt.")
    rows = [[_cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
    table = _table_from_rows(rows)
    table.template, table.template_name = template, name
    return table


def _embedded_template(wb: Any) -> tuple[dict[str, Any] | None, str]:
    if TEMPLATE_SHEET not in wb.sheetnames:
        return None, ""
    ws = wb[TEMPLATE_SHEET]
    if ws["A1"].value != TEMPLATE_MARKER:
        return None, ""
    chunks = [str(r[0])[1:] for r in ws.iter_rows(min_row=2, max_col=1, values_only=True) if r[0]]
    try:
        data = json.loads("".join(chunks))
        Label.from_dict(data)
    except (ValueError, KeyError, TypeError):
        return None, ""
    return data, str(ws["B1"].value or "")


def _read_csv(path: Path) -> Table:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    first = next((line for line in text.splitlines() if line.strip()), "")
    delimiter = max(";,\t", key=lambda d: (first.count(d), d == ";"))
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    return _table_from_rows(rows)


def _table_from_rows(rows: list[list[str]]) -> Table:
    header_at = next((i for i, row in enumerate(rows) if any(c.strip() for c in row)), None)
    if header_at is None:
        raise TableError("Die Tabelle ist leer.")

    columns: list[str] = []
    positions: list[int] = []
    problems: list[str] = []
    seen: set[str] = set()
    for i, cell in enumerate(rows[header_at]):
        name = cell.strip()
        if not name:
            continue
        if merge.key(name) in seen:
            problems.append(f"Die Spalte „{name}“ gibt es doppelt, verwendet wird die erste.")
            continue
        seen.add(merge.key(name))
        columns.append(name)
        positions.append(i)

    data: list[dict[str, str]] = []
    lines: list[int] = []
    for line, row in enumerate(rows[header_at + 1 :], start=header_at + 2):
        record = {name: (row[i] if i < len(row) else "") for name, i in zip(columns, positions, strict=True)}
        if any(v.strip() for v in record.values()):
            data.append(record)
            lines.append(line)
    return Table(columns, data, lines, warnings=problems)
