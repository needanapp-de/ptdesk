from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from ptdesk.render import merge
from ptdesk.render.label import BarcodeElement, Label, TextElement
from ptdesk.table import TableError, read_table, write_template


def label() -> Label:
    return Label(
        elements=[
            TextElement(text="{Name}\n{Raum}"),
            BarcodeElement(data="{Inventar}", symbology="code128"),
            TextElement(text="={Hinweis}"),
        ],
        sample={"Name": "Drucker", "Raum": "Büro 1", "Inventar": "A-1", "Hinweis": "=nicht rechnen"},
    )


def test_xlsx_template_roundtrip(tmp_path: Path):
    path = tmp_path / "liste.xlsx"
    write_template(path, label(), "inventar.json")
    table = read_table(path)
    assert table.columns == ["Name", "Raum", "Inventar", "Hinweis", "Anzahl"]
    assert table.rows == [
        {"Name": "Drucker", "Raum": "Büro 1", "Inventar": "A-1", "Hinweis": "=nicht rechnen", "Anzahl": "1"}
    ]
    assert table.lines == [2]
    assert table.template_name == "inventar.json"
    assert Label.from_dict(table.template).sample == label().sample

    wb = load_workbook(path)
    ws = wb["Etiketten"]
    assert ws.column_dimensions["A"].number_format == "@"  # "+49…" and leading zeros stay text
    assert "Verwendet in: Code 128" in ws["C1"].comment.text
    assert [s.sheet_state for s in wb.worksheets] == ["visible", "visible", "veryHidden"]


def test_filled_xlsx_with_typed_cells(tmp_path: Path):
    path = tmp_path / "liste.xlsx"
    write_template(path, label())
    wb = load_workbook(path)
    ws = wb["Etiketten"]
    ws.append(["Monitor", "Lager\nRegal 3", 12345678, 1.5, 2])
    ws.append([None, None, None, None, None])  # empty rows are skipped
    ws.append(["Kabel", "Keller", "K-9", datetime(2026, 12, 24, 18, 30), None])  # noqa: DTZ001
    ws.append(["Kiste", "Keller", "K-10", date(2026, 1, 2), True])
    wb.save(path)

    table = read_table(path)
    assert table.lines == [2, 3, 5, 6]
    assert table.rows[1] == {"Name": "Monitor", "Raum": "Lager\nRegal 3", "Inventar": "12345678", "Hinweis": "1,5", "Anzahl": "2"}
    assert table.rows[2]["Hinweis"] == "24.12.2026 18:30"
    assert table.rows[3]["Hinweis"] == "02.01.2026"

    merged = merge.merge_rows(label(), table.rows, table.lines, example=merge.example_values(label()))
    assert merged[0].is_example
    assert merged[1].ok and merged[1].copies == 2
    assert merged[1].label.elements[0].text == "Monitor\nLager\nRegal 3"
    assert merged[1].label.elements[1].data == "12345678"
    assert merged[2].ok and merged[2].copies == 1
    assert not merged[3].ok  # "ja" is no number of copies


def test_csv_template_and_foreign_csv(tmp_path: Path):
    path = tmp_path / "liste.csv"
    write_template(path, label())
    assert path.read_bytes().startswith(b"\xef\xbb\xbfName;Raum;")
    assert read_table(path).columns[-1] == "Anzahl"

    other = tmp_path / "excel.csv"
    other.write_bytes("Name,Straße\n\nMüller,\"Hauptstr. 1, Hinterhaus\"\n".encode("cp1252"))
    table = read_table(other)
    assert table.columns == ["Name", "Straße"]
    assert table.rows == [{"Name": "Müller", "Straße": "Hauptstr. 1, Hinterhaus"}]
    assert table.lines == [3]


def test_foreign_xlsx_header_not_in_first_row(tmp_path: Path):
    path = tmp_path / "fremd.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append([])
    ws.append(["Name", "Name", None, "Nr"])
    ws.append(["A", "B", "ignoriert", 7])
    wb.save(path)
    table = read_table(path)
    assert table.columns == ["Name", "Nr"]
    assert table.rows == [{"Name": "A", "Nr": "7"}]
    assert table.lines == [3]
    assert table.template is None
    assert table.warnings


def test_errors(tmp_path: Path):
    with pytest.raises(TableError):
        write_template(tmp_path / "x.xlsx", Label(elements=[TextElement(text="fest")]))
    with pytest.raises(TableError):
        read_table(tmp_path / "x.ods")
    broken = tmp_path / "kaputt.xlsx"
    broken.write_text("kein zip")
    with pytest.raises(TableError):
        read_table(broken)
    empty = tmp_path / "leer.csv"
    empty.write_text("\n\n")
    with pytest.raises(TableError):
        read_table(empty)


def test_cli_render_from_table(tmp_path: Path, capsys):
    import json

    from ptdesk.cli import main, parse_rows

    assert parse_rows("3-5, 8") == {3, 4, 5, 8}
    template = tmp_path / "vorlage.json"
    template.write_text(json.dumps(label().to_dict()), encoding="utf-8")
    with pytest.raises(SystemExit) as done:
        main(["table", str(template)])
    assert done.value.code == 0
    table = tmp_path / "vorlage-tabelle.xlsx"
    wb = load_workbook(table)
    wb["Etiketten"].append(["Monitor", "Büro", "M-1", "x", 2])
    wb["Etiketten"].append(["Kabel", "Lager", "", "x", 1])
    wb.save(table)

    out = tmp_path / "e.png"
    with pytest.raises(SystemExit) as done:
        main(["render", str(table), "-o", str(out)])  # row 4 has an empty barcode
    assert done.value.code == 1
    with pytest.raises(SystemExit) as done:
        main(["render", str(table), "--skip-errors", "-o", str(out)])
    assert done.value.code == 0
    assert (tmp_path / "e-003.png").exists() and not (tmp_path / "e-004.png").exists()
    err = capsys.readouterr().err
    assert "Zeile 2: Beispielzeile" in err and "Zeile 4: Barcode: Barcode ist leer" in err
